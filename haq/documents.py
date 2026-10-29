"""Local text extraction for uploaded notice documents.

Only PDF, DOCX, and TXT files are accepted. Files are read in-process and are
never uploaded to an external service. Scanned PDF pages are reported as
having no extractable text; this module does not perform OCR.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 200_000
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")
_KIND_LABELS = {".pdf": "PDF", ".docx": "DOCX", ".txt": "plain text"}


class DocumentError(ValueError):
    """A document could not be read; the message is safe to show to users."""


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    kind: str
    pages: int | None
    warnings: tuple[str, ...]


def _coerce_path(file_value) -> Path:
    if isinstance(file_value, (list, tuple)):
        file_value = file_value[0] if file_value else None
    if file_value in (None, ""):
        raise DocumentError("Choose a PDF, DOCX, or TXT file first.")
    return Path(str(file_value))


def _validate(file_value) -> tuple[Path, str]:
    path = _coerce_path(file_value)
    suffix = path.suffix.lower()
    if suffix not in _KIND_LABELS:
        raise DocumentError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    if not path.is_file():
        raise DocumentError("The uploaded file could not be found. Please upload it again.")

    size = path.stat().st_size
    if size == 0:
        raise DocumentError("The uploaded file is empty.")
    if size > MAX_FILE_BYTES:
        limit_mb = MAX_FILE_BYTES // (1024 * 1024)
        raise DocumentError(f"The file is larger than {limit_mb} MB. Upload a smaller file.")
    return path, suffix


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.split("\n")]
    collapsed: list[str] = []
    blank_run = 0
    for line in lines:
        if line:
            collapsed.append(line)
            blank_run = 0
            continue
        blank_run += 1
        if blank_run == 1:
            collapsed.append("")
    return "\n".join(collapsed).strip()


def _read_txt(path: Path) -> tuple[str, int | None, tuple[str, ...]]:
    data = path.read_bytes()
    encodings = ["utf-8-sig"]
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings.append("utf-16")
    for encoding in encodings:
        try:
            return data.decode(encoding), None, ()
        except UnicodeError:
            continue
    raise DocumentError(
        "The TXT file is not valid UTF-8 or UTF-16 text. "
        "Re-save it as UTF-8 and upload it again."
    )


def _read_docx(path: Path) -> tuple[str, int | None, tuple[str, ...]]:
    try:
        from docx import Document
    except ImportError:  # pragma: no cover - dependency is pinned in requirements
        raise DocumentError("DOCX support is not installed on this system.") from None

    try:
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
    except Exception:
        raise DocumentError(
            "This DOCX could not be read. It may be corrupt, password-protected, "
            "or not a Word document."
        ) from None
    return "\n".join(parts), None, ()


def _read_pdf(path: Path) -> tuple[str, int | None, tuple[str, ...]]:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - dependency is pinned in requirements
        raise DocumentError("PDF support is not installed on this system.") from None

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                unlocked = reader.decrypt("")
            except Exception:
                unlocked = 0
            if not unlocked:
                raise DocumentError(
                    "This PDF is password-protected. Remove the password and upload it again."
                )

        pages = reader.pages
        page_count = len(pages)
        if page_count == 0:
            raise DocumentError("This PDF has no pages.")

        chunks: list[str] = []
        empty_pages = 0
        for page in pages:
            try:
                page_text = page.extract_text() or ""
            except Exception:
                page_text = ""
            if not page_text.strip():
                empty_pages += 1
            chunks.append(page_text)
    except DocumentError:
        raise
    except Exception:
        raise DocumentError(
            "This PDF could not be read. It may be corrupt or use an unsupported format."
        ) from None

    text = "\n".join(chunks)
    if not text.strip():
        raise DocumentError(
            "This PDF has no extractable text. It may contain scanned images; "
            "OCR is not supported."
        )

    warnings: tuple[str, ...] = ()
    if empty_pages:
        warnings = (
            f"{empty_pages} of {page_count} PDF page(s) had no extractable text; "
            "they may be scanned images. OCR is not supported.",
        )
    return text, page_count, warnings


_READERS = {".txt": _read_txt, ".docx": _read_docx, ".pdf": _read_pdf}


def read_notice_document(file_value) -> ExtractedDocument:
    """Extract readable text from a supported notice file.

    Raises DocumentError with a user-safe message when the file is unsupported,
    empty, too large, encrypted, malformed, or has no extractable text.
    """

    path, suffix = _validate(file_value)
    text, pages, warnings = _READERS[suffix](path)
    normalised = _normalise(text)
    if not normalised:
        raise DocumentError("No readable text was found in this file.")
    if len(normalised) > MAX_TEXT_CHARS:
        raise DocumentError(
            f"The extracted text is longer than {MAX_TEXT_CHARS:,} characters. "
            "Upload a shorter notice."
        )
    return ExtractedDocument(
        text=normalised,
        kind=_KIND_LABELS[suffix],
        pages=pages,
        warnings=warnings,
    )
