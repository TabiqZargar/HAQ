import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from haq.documents import DocumentError, read_notice_document
from haq.notice import extract_notice_facts


def _write_pdf(path, page_texts):
    writer = PdfWriter()
    for text in page_texts:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        if text:
            escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1"))
            page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(str(path))
    return path


def _write_docx(path, paragraphs, table_rows=()):
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    document.save(str(path))
    return path


def test_reads_plain_text_and_preserves_explicit_facts(tmp_path) -> None:
    path = tmp_path / "notice.txt"
    path.write_text(
        "From: Rakesh Sharma\r\nYou must vacate on or before 30/11/2026.\r\n"
        "Payment of Rs. 45,000 is due within 15 days.\r\n",
        encoding="utf-8",
    )

    document = read_notice_document(path)

    assert document.kind == "plain text"
    assert document.pages is None
    assert document.warnings == ()
    assert "\r" not in document.text
    facts = extract_notice_facts(document.text)
    assert "30/11/2026" in facts.dates
    assert "Rs. 45,000" in facts.amounts
    assert "Rakesh Sharma" in facts.parties
    assert any("within 15 days" in phrase for phrase in facts.deadline_language)


def test_reads_docx_paragraphs_and_tables(tmp_path) -> None:
    path = _write_docx(
        tmp_path / "notice.docx",
        ["NOTICE NO. LD/2026/045", "You must vacate on or before 30 November 2026."],
        table_rows=[("Amount", "Rs. 10,000")],
    )

    document = read_notice_document(path)

    assert document.kind == "DOCX"
    assert "LD/2026/045" in document.text
    assert "30 November 2026" in document.text
    assert "Amount | Rs. 10,000" in document.text


def test_reads_text_from_pdf(tmp_path) -> None:
    path = _write_pdf(
        tmp_path / "notice.pdf",
        ["You must vacate on or before 30/11/2026.", "Payment of Rs. 45,000 is due."],
    )

    document = read_notice_document(path)

    assert document.kind == "PDF"
    assert document.pages == 2
    assert document.warnings == ()
    assert "30/11/2026" in document.text
    assert "Rs. 45,000" in document.text


def test_pdf_with_no_extractable_text_reports_no_ocr(tmp_path) -> None:
    path = _write_pdf(tmp_path / "scanned.pdf", ["", ""])

    with pytest.raises(DocumentError, match="OCR is not supported"):
        read_notice_document(path)


def test_pdf_with_some_empty_pages_returns_a_warning(tmp_path) -> None:
    path = _write_pdf(tmp_path / "mixed.pdf", ["Readable text.", ""])

    document = read_notice_document(path)

    assert "Readable text." in document.text
    assert len(document.warnings) == 1
    assert "1 of 2" in document.warnings[0]
    assert "OCR is not supported" in document.warnings[0]


def test_encrypted_pdf_is_rejected(tmp_path) -> None:
    path = tmp_path / "locked.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret")
    writer.write(str(path))

    with pytest.raises(DocumentError, match="password-protected"):
        read_notice_document(path)


def test_unsupported_extension_is_rejected(tmp_path) -> None:
    path = tmp_path / "notice.rtf"
    path.write_text("{\\rtf1 notice}", encoding="utf-8")

    with pytest.raises(DocumentError, match="Unsupported file type"):
        read_notice_document(path)


def test_empty_file_is_rejected(tmp_path) -> None:
    path = tmp_path / "empty.txt"
    path.write_bytes(b"")

    with pytest.raises(DocumentError, match="empty"):
        read_notice_document(path)


def test_oversized_file_is_rejected(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("haq.documents.MAX_FILE_BYTES", 1024)
    path = tmp_path / "large.txt"
    path.write_bytes(b"x" * 1025)

    with pytest.raises(DocumentError, match="larger than"):
        read_notice_document(path)


def test_malformed_pdf_and_docx_are_rejected_gracefully(tmp_path) -> None:
    pdf = tmp_path / "broken.pdf"
    pdf.write_bytes(b"this is not a pdf")
    docx = tmp_path / "broken.docx"
    docx.write_bytes(b"this is not a zip archive")

    with pytest.raises(DocumentError):
        read_notice_document(pdf)
    with pytest.raises(DocumentError):
        read_notice_document(docx)


def test_invalid_text_encoding_is_rejected(tmp_path) -> None:
    path = tmp_path / "notice.txt"
    path.write_bytes(b"valid prefix \xff\xfe\xff invalid bytes")

    with pytest.raises(DocumentError, match="UTF-8 or UTF-16"):
        read_notice_document(path)


def test_reads_utf16_text_with_a_byte_order_mark(tmp_path) -> None:
    path = tmp_path / "notice.txt"
    path.write_bytes("You must vacate on or before 30/11/2026.".encode("utf-16"))

    document = read_notice_document(path)

    assert "30/11/2026" in document.text


def test_docx_without_text_is_rejected(tmp_path) -> None:
    path = _write_docx(tmp_path / "blank.docx", ["", "   "])

    with pytest.raises(DocumentError, match="No readable text"):
        read_notice_document(path)
